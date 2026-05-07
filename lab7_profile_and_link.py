import csv
import math
import time
from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
import requests
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


OUTPUT_DIR = Path("output_lab7")
SHOW_PLOTS = False
DEM_DATASET = "srtm90m"
DEM_SAMPLES = 41
FREQUENCY_MHZ = 5500.0
CHANNEL_WIDTH_MHZ = 80
NOISE_FREE_FLOOR_DBM = -100.0
MAX_PHY_MBPS = 650.0
SYSTEM_LOSS_DB = 1.5


@dataclass(frozen=True)
class Equipment:
    name: str
    tx_power_max_dbm: float
    tx_power_used_dbm: float
    antenna_gain_dbi: float
    max_range_km: float
    max_ideal_mbps: float
    frequency_mhz: float


@dataclass(frozen=True)
class Site:
    name: str
    lat: float
    lon: float
    placement_note: str


@dataclass(frozen=True)
class LinkConfig:
    name: str
    tx_site: Site
    rx_site: Site
    tx_height_m: float
    rx_height_m: float
    tx_equipment: Equipment
    rx_equipment: Equipment
    obstacle_type: str
    note: str


@dataclass(frozen=True)
class RadioScenario:
    name: str
    rain_mm_h: float
    noise_floor_dbm: float


@dataclass(frozen=True)
class SensitivityPoint:
    label: str
    sensitivity_dbm: float
    spectral_efficiency: float


SENSITIVITY_TABLE = [
    SensitivityPoint("MCS0 / BPSK 1/2", -96.0, 0.5),
    SensitivityPoint("MCS1 / QPSK 1/2", -95.0, 1.0),
    SensitivityPoint("MCS2 / QPSK 3/4", -92.0, 1.5),
    SensitivityPoint("MCS3 / 16QAM 1/2", -90.0, 2.0),
    SensitivityPoint("MCS4 / 16QAM 3/4", -86.0, 3.0),
    SensitivityPoint("MCS5 / 64QAM 2/3", -83.0, 4.0),
    SensitivityPoint("MCS6 / 64QAM 3/4", -77.0, 4.5),
    SensitivityPoint("MCS7 / 64QAM 5/6", -74.0, 5.0),
    SensitivityPoint("MCS8 / 256QAM 3/4", -69.0, 6.0),
    SensitivityPoint("MCS9 / 256QAM 5/6", -65.0, 6.6667),
]


RADIO_SCENARIOS = [
    RadioScenario("Noise Free, без опадів", 0.0, -100.0),
    RadioScenario("Noise Free, дощ 5 мм/год", 5.0, -100.0),
    RadioScenario("Rural, без опадів", 0.0, -96.0),
    RadioScenario("Urban, без опадів", 0.0, -89.0),
]


LITEBEAM_5AC = Equipment(
    name="Ubiquiti airMAX LiteBeam 5AC Gen2",
    tx_power_max_dbm=25.0,
    tx_power_used_dbm=8.0,
    antenna_gain_dbi=23.0,
    max_range_km=30.0,
    max_ideal_mbps=450.0,
    frequency_mhz=FREQUENCY_MHZ,
)


LITEAP_GPS = Equipment(
    name="Ubiquiti airMAX Lite AP GPS",
    tx_power_max_dbm=25.0,
    tx_power_used_dbm=10.0,
    antenna_gain_dbi=17.0,
    max_range_km=30.0,
    max_ideal_mbps=450.0,
    frequency_mhz=FREQUENCY_MHZ,
)


MAIN_BUILDING = Site(
    name="Головний корпус ХАІ",
    lat=50.0428180,
    lon=36.2851866,
    placement_note="дах головного корпусу ХАІ",
)


DORM_10 = Site(
    name="Гуртожиток ХАІ №10",
    lat=50.0413728,
    lon=36.2950078,
    placement_note="будівля гуртожитку №10",
)


DORM_11 = Site(
    name="Гуртожиток ХАІ №11",
    lat=50.0406940,
    lon=36.2949133,
    placement_note="будівля гуртожитку №11",
)


DORM_12 = Site(
    name="Гуртожиток ХАІ №12",
    lat=50.0397665,
    lon=36.2949418,
    placement_note="будівля гуртожитку №12",
)


PTP_NLOS = LinkConfig(
    name="PtP nLOS: Головний корпус ХАІ - гуртожиток №10",
    tx_site=MAIN_BUILDING,
    rx_site=DORM_10,
    tx_height_m=15.0,
    rx_height_m=3.0,
    tx_equipment=LITEBEAM_5AC,
    rx_equipment=LITEBEAM_5AC,
    obstacle_type="хвилястий рельєф і щільна міська забудова вздовж напрямку на гуртожитки",
    note="Базова висота приймача біля рівня нижніх поверхів/прибудови. Пряма видимість є, але зона Френеля перетинається.",
)


PTP_LOS = LinkConfig(
    name="PtP LOS: Головний корпус ХАІ - гуртожиток №10",
    tx_site=MAIN_BUILDING,
    rx_site=DORM_10,
    tx_height_m=15.0,
    rx_height_m=12.0,
    tx_equipment=LITEBEAM_5AC,
    rx_equipment=LITEBEAM_5AC,
    obstacle_type="той самий рельєф; антена піднята на рівень верхніх поверхів",
    note="Антену приймача піднято на рівень 4-го поверху, щоб вивести трасу з режиму nLOS до LOS.",
)


PTMP_LINKS = [
    LinkConfig(
        name="PtMP AP1 -> Station1 (гуртожиток №10)",
        tx_site=MAIN_BUILDING,
        rx_site=DORM_10,
        tx_height_m=18.0,
        rx_height_m=12.0,
        tx_equipment=LITEAP_GPS,
        rx_equipment=LITEBEAM_5AC,
        obstacle_type="міська забудова та окремі дерева, що залишаються нижче першої зони Френеля",
        note="Центральна точка доступу на головному корпусі, абонентська станція на даху гуртожитку №10.",
    ),
    LinkConfig(
        name="PtMP AP1 -> Station2 (гуртожиток №11)",
        tx_site=MAIN_BUILDING,
        rx_site=DORM_11,
        tx_height_m=18.0,
        rx_height_m=12.0,
        tx_equipment=LITEAP_GPS,
        rx_equipment=LITEBEAM_5AC,
        obstacle_type="міська забудова та окремі дерева, що залишаються нижче першої зони Френеля",
        note="Центральна точка доступу на головному корпусі, абонентська станція на даху гуртожитку №11.",
    ),
    LinkConfig(
        name="PtMP AP1 -> Station3 (гуртожиток №12)",
        tx_site=MAIN_BUILDING,
        rx_site=DORM_12,
        tx_height_m=18.0,
        rx_height_m=12.0,
        tx_equipment=LITEAP_GPS,
        rx_equipment=LITEBEAM_5AC,
        obstacle_type="міська забудова та окремі дерева, що залишаються нижче першої зони Френеля",
        note="Центральна точка доступу на головному корпусі, абонентська станція на даху гуртожитку №12.",
    ),
]


def haversine_distance_m(site_a: Site, site_b: Site) -> float:
    radius_m = 6_371_000.0
    lat1 = math.radians(site_a.lat)
    lat2 = math.radians(site_b.lat)
    dlat = math.radians(site_b.lat - site_a.lat)
    dlon = math.radians(site_b.lon - site_a.lon)
    term = math.sin(dlat / 2.0) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2.0) ** 2
    return radius_m * 2.0 * math.atan2(math.sqrt(term), math.sqrt(1.0 - term))


def fetch_profile(link: LinkConfig, *, samples: int = DEM_SAMPLES, dataset: str = DEM_DATASET) -> tuple[list[float], list[float]]:
    url = f"https://api.opentopodata.org/v1/{dataset}"
    params = {
        "locations": f"{link.tx_site.lat},{link.tx_site.lon}|{link.rx_site.lat},{link.rx_site.lon}",
        "samples": samples,
    }

    last_error = None
    for _ in range(3):
        try:
            response = requests.get(url, params=params, timeout=60)
            data = response.json()
            if response.status_code == 200 and "results" in data:
                elevations = [float(item["elevation"]) for item in data["results"]]
                distance_m = haversine_distance_m(link.tx_site, link.rx_site)
                distances = [distance_m * idx / (samples - 1) for idx in range(samples)]
                time.sleep(1.1)
                return distances, elevations
            last_error = RuntimeError(f"{response.status_code}: {data}")
        except Exception as exc:  # pragma: no cover - network edge
            last_error = exc
        time.sleep(2.0)

    raise RuntimeError(f"Не вдалося отримати профіль висот для {link.name}: {last_error}")


def fresnel_radius_m(distance_from_tx_m: float, total_distance_m: float, frequency_mhz: float) -> float:
    if distance_from_tx_m <= 0.0 or distance_from_tx_m >= total_distance_m:
        return 0.0
    wavelength_m = 300.0 / frequency_mhz
    return math.sqrt(wavelength_m * distance_from_tx_m * (total_distance_m - distance_from_tx_m) / total_distance_m)


def free_space_path_loss_db(distance_km: float, frequency_mhz: float) -> float:
    return 32.44 + 20.0 * math.log10(distance_km) + 20.0 * math.log10(frequency_mhz)


def rain_loss_db(distance_km: float, rain_mm_h: float) -> float:
    if rain_mm_h <= 0.0:
        return 0.0
    specific_attenuation_db_per_km = 0.02 * (rain_mm_h / 5.0)
    return specific_attenuation_db_per_km * distance_km


def classify_link(clearance_m: list[float], fresnel_m: list[float]) -> str:
    min_los = min(clearance_m[1:-1])
    min_fresnel = min(clearance_m[idx] - fresnel_m[idx] for idx in range(1, len(clearance_m) - 1))
    if min_fresnel > 0.0:
        return "LOS"
    if min_los > 0.0:
        return "nLOS"
    return "NLOS"


def mcs_and_capacity(signal_dbm: float, noise_floor_dbm: float) -> tuple[str, float]:
    effective_signal = signal_dbm - max(0.0, noise_floor_dbm - NOISE_FREE_FLOOR_DBM)
    chosen = SENSITIVITY_TABLE[0]
    for point in SENSITIVITY_TABLE:
        if effective_signal >= point.sensitivity_dbm:
            chosen = point
    capacity_mbps = MAX_PHY_MBPS * chosen.spectral_efficiency / SENSITIVITY_TABLE[-1].spectral_efficiency
    return chosen.label, round(capacity_mbps, 1)


def analyze_link(link: LinkConfig) -> dict:
    distances_m, elevations_m = fetch_profile(link)
    total_distance_m = distances_m[-1]
    tx_abs_m = elevations_m[0] + link.tx_height_m
    rx_abs_m = elevations_m[-1] + link.rx_height_m

    los_line_m = [
        tx_abs_m + (rx_abs_m - tx_abs_m) * (distance / total_distance_m)
        for distance in distances_m
    ]
    fresnel_m = [fresnel_radius_m(distance, total_distance_m, FREQUENCY_MHZ) for distance in distances_m]
    clearance_m = [los_line_m[idx] - elevations_m[idx] for idx in range(len(distances_m))]
    link_type = classify_link(clearance_m, fresnel_m)
    max_obstacle_idx = min(range(1, len(clearance_m) - 1), key=lambda idx: clearance_m[idx])
    midpoint_idx = len(distances_m) // 2

    obstacle_distance_m = distances_m[max_obstacle_idx]
    los_clearance_m = clearance_m[max_obstacle_idx]
    fresnel_at_obstacle_m = fresnel_m[max_obstacle_idx]
    fresnel_mid_m = fresnel_m[midpoint_idx]
    fresnel_intrusion_m = max(0.0, fresnel_at_obstacle_m - los_clearance_m)
    fresnel_penalty_db = 0.0
    if fresnel_at_obstacle_m > 0.0 and fresnel_intrusion_m > 0.0:
        fresnel_penalty_db = min(8.0, 4.0 * fresnel_intrusion_m / fresnel_at_obstacle_m)

    distance_km = total_distance_m / 1000.0
    base_signal_dbm = (
        link.tx_equipment.tx_power_used_dbm
        + link.tx_equipment.antenna_gain_dbi
        + link.rx_equipment.antenna_gain_dbi
        - free_space_path_loss_db(distance_km, FREQUENCY_MHZ)
        - SYSTEM_LOSS_DB
        - fresnel_penalty_db
    )

    radio_rows = []
    for scenario in RADIO_SCENARIOS:
        signal_dbm = base_signal_dbm - rain_loss_db(distance_km, scenario.rain_mm_h)
        mcs_label, capacity_mbps = mcs_and_capacity(signal_dbm, scenario.noise_floor_dbm)
        radio_rows.append(
            {
                "scenario": scenario.name,
                "signal_dbm": round(signal_dbm, 1),
                "capacity_mbps": capacity_mbps,
                "mcs": mcs_label,
                "noise_floor_dbm": scenario.noise_floor_dbm,
            }
        )

    return {
        "link": link,
        "distances_m": distances_m,
        "elevations_m": elevations_m,
        "los_line_m": los_line_m,
        "fresnel_m": fresnel_m,
        "clearance_m": clearance_m,
        "distance_m": total_distance_m,
        "distance_km": distance_km,
        "link_type": link_type,
        "max_obstacle_idx": max_obstacle_idx,
        "max_obstacle_distance_m": obstacle_distance_m,
        "los_clearance_m": los_clearance_m,
        "fresnel_at_obstacle_m": fresnel_at_obstacle_m,
        "fresnel_mid_m": fresnel_mid_m,
        "fresnel_intrusion_m": fresnel_intrusion_m,
        "fresnel_penalty_db": round(fresnel_penalty_db, 2),
        "base_signal_dbm": round(base_signal_dbm, 1),
        "radio_rows": radio_rows,
    }


def plot_profile(ax, result: dict, title: str) -> None:
    distances = [value / 1000.0 for value in result["distances_m"]]
    terrain = result["elevations_m"]
    los_line = result["los_line_m"]
    fresnel = result["fresnel_m"]
    upper = [los_line[idx] + fresnel[idx] for idx in range(len(los_line))]
    lower = [los_line[idx] - fresnel[idx] for idx in range(len(los_line))]
    idx = result["max_obstacle_idx"]

    ax.fill_between(distances, terrain, color="#c7b299", alpha=0.75, label="Профіль траси")
    ax.plot(distances, los_line, color="#1f77b4", linewidth=1.4, label="Лінія прямої видимості")
    ax.plot(distances, upper, color="#2ca02c", linewidth=0.9, linestyle="--", label="1-а зона Френеля")
    ax.plot(distances, lower, color="#2ca02c", linewidth=0.9, linestyle="--")
    ax.scatter([distances[idx]], [terrain[idx]], color="#d62728", s=22, zorder=3, label="Макс. перешкода")
    ax.set_title(title, fontsize=10)
    ax.set_xlabel("Відстань, км")
    ax.set_ylabel("Висота, м")
    ax.grid(True, linestyle="--", linewidth=0.35, alpha=0.6)


def save_profile_figures(ptp_results: list[dict], ptmp_results: list[dict]) -> dict[str, Path]:
    figures = {}

    fig, axes = plt.subplots(2, 1, figsize=(19 / 2.54, 14 / 2.54), sharex=False)
    plot_profile(axes[0], ptp_results[0], "PtP nLOS: головний корпус ХАІ - гуртожиток №10")
    plot_profile(axes[1], ptp_results[1], "PtP LOS: головний корпус ХАІ - гуртожиток №10")
    handles, labels = axes[0].get_legend_handles_labels()
    axes[0].legend(handles, labels, loc="upper right", fontsize=8)
    fig.tight_layout()
    figures["ptp_profiles"] = OUTPUT_DIR / "ptp_profiles.png"
    fig.savefig(figures["ptp_profiles"], dpi=300)
    plt.close(fig)

    fig, axes = plt.subplots(3, 1, figsize=(19 / 2.54, 19 / 2.54), sharex=False)
    titles = [
        "PtMP AP1 -> Station1 (гуртожиток №10)",
        "PtMP AP1 -> Station2 (гуртожиток №11)",
        "PtMP AP1 -> Station3 (гуртожиток №12)",
    ]
    for axis, result, title in zip(axes, ptmp_results, titles):
        plot_profile(axis, result, title)
    handles, labels = axes[0].get_legend_handles_labels()
    axes[0].legend(handles, labels, loc="upper right", fontsize=7)
    fig.tight_layout()
    figures["ptmp_profiles"] = OUTPUT_DIR / "ptmp_profiles.png"
    fig.savefig(figures["ptmp_profiles"], dpi=300)
    plt.close(fig)

    figures["ptp_scheme"] = OUTPUT_DIR / "ptp_scheme.png"
    build_scheme_figure([ptp_results[0]["link"].tx_site, ptp_results[0]["link"].rx_site], [(0, 1)], figures["ptp_scheme"])

    figures["ptmp_scheme"] = OUTPUT_DIR / "ptmp_scheme.png"
    build_scheme_figure(
        [MAIN_BUILDING, DORM_10, DORM_11, DORM_12],
        [(0, 1), (0, 2), (0, 3)],
        figures["ptmp_scheme"],
    )

    return figures


def build_scheme_figure(sites: list[Site], edges: list[tuple[int, int]], output_path: Path) -> None:
    mean_lat_rad = math.radians(sum(site.lat for site in sites) / len(sites))
    origin_lat = sites[0].lat
    origin_lon = sites[0].lon

    xy_points = []
    for site in sites:
        x = (site.lon - origin_lon) * 111_320.0 * math.cos(mean_lat_rad)
        y = (site.lat - origin_lat) * 110_540.0
        xy_points.append((x, y))

    fig, ax = plt.subplots(figsize=(16 / 2.54, 10 / 2.54))
    for left_idx, right_idx in edges:
        x_values = [xy_points[left_idx][0], xy_points[right_idx][0]]
        y_values = [xy_points[left_idx][1], xy_points[right_idx][1]]
        ax.plot(x_values, y_values, color="#1f77b4", linewidth=1.2)

    for site, (x_coord, y_coord) in zip(sites, xy_points):
        ax.scatter(x_coord, y_coord, color="#d62728", s=28)
        ax.text(x_coord + 7.0, y_coord + 7.0, site.name, fontsize=8)

    ax.set_xlabel("Відносна координата X, м")
    ax.set_ylabel("Відносна координата Y, м")
    ax.set_title("Розташування вузлів")
    ax.grid(True, linestyle="--", linewidth=0.35, alpha=0.6)
    fig.tight_layout()
    fig.savefig(output_path, dpi=300)
    plt.close(fig)


def save_csv_outputs(results: list[dict]) -> None:
    summary_path = OUTPUT_DIR / "link_summary.csv"
    scenario_path = OUTPUT_DIR / "radio_scenarios.csv"
    profile_path = OUTPUT_DIR / "profiles.csv"

    with summary_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(
            [
                "link_name",
                "link_type",
                "distance_m",
                "tx_height_m",
                "rx_height_m",
                "max_obstacle_distance_m",
                "los_clearance_m",
                "fresnel_at_obstacle_m",
                "fresnel_mid_m",
                "fresnel_intrusion_m",
            ]
        )
        for result in results:
            link = result["link"]
            writer.writerow(
                [
                    link.name,
                    result["link_type"],
                    round(result["distance_m"], 1),
                    link.tx_height_m,
                    link.rx_height_m,
                    round(result["max_obstacle_distance_m"], 1),
                    round(result["los_clearance_m"], 2),
                    round(result["fresnel_at_obstacle_m"], 2),
                    round(result["fresnel_mid_m"], 2),
                    round(result["fresnel_intrusion_m"], 2),
                ]
            )

    with scenario_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["link_name", "scenario", "signal_dbm", "capacity_mbps", "mcs", "noise_floor_dbm"])
        for result in results:
            for row in result["radio_rows"]:
                writer.writerow(
                    [
                        result["link"].name,
                        row["scenario"],
                        row["signal_dbm"],
                        row["capacity_mbps"],
                        row["mcs"],
                        row["noise_floor_dbm"],
                    ]
                )

    with profile_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["link_name", "distance_m", "terrain_m", "los_line_m", "fresnel_m", "clearance_m"])
        for result in results:
            for idx in range(len(result["distances_m"])):
                writer.writerow(
                    [
                        result["link"].name,
                        round(result["distances_m"][idx], 2),
                        round(result["elevations_m"][idx], 2),
                        round(result["los_line_m"][idx], 2),
                        round(result["fresnel_m"][idx], 2),
                        round(result["clearance_m"][idx], 2),
                    ]
                )


def format_link_result(result: dict) -> list[str]:
    link = result["link"]
    lines = [
        f"### {link.name}",
        f"- Довжина радіолінка: {result['distance_m']:.1f} м ({result['distance_km']:.3f} км)",
        f"- Тип траси: {result['link_type']}",
        f"- Розташування антен: передавач {link.tx_height_m:.1f} м, приймач {link.rx_height_m:.1f} м",
        f"- Тип основних перешкод: {link.obstacle_type}",
        f"- Максимальна перешкода на відстані {result['max_obstacle_distance_m']:.1f} м від передавача",
        f"- Просвіт LOS у точці максимальної перешкоди: {result['los_clearance_m']:.2f} м",
        f"- Перша зона Френеля у точці максимальної перешкоди: {result['fresnel_at_obstacle_m']:.2f} м",
        f"- Перша зона Френеля на середині траси: {result['fresnel_mid_m']:.2f} м",
        f"- Заходження в 1-у зону Френеля: {result['fresnel_intrusion_m']:.2f} м",
        f"- Примітка: {link.note}",
        "",
        "| Сценарій | Expected Signal, dBm | Capacity, Мбіт/с | Модуляція |",
        "| --- | ---: | ---: | --- |",
    ]
    for row in result["radio_rows"]:
        lines.append(
            f"| {row['scenario']} | {row['signal_dbm']:.1f} | {row['capacity_mbps']:.1f} | {row['mcs']} |"
        )
    lines.append("")
    return lines


def build_markdown_report(results_ptp: list[dict], results_ptmp: list[dict], figures: dict[str, Path]) -> None:
    report_path = OUTPUT_DIR / "lab7_report.md"
    lines = [
        "# Лабораторна робота 7",
        "",
        "## Тема",
        "Побудова профілю траси та визначення її основних параметрів.",
        "",
        "## Мета",
        "Закріпити на практиці поняття зони Френеля та профілю траси; побудувати радіолінії PtP і PtMP на території України та визначити їх основні параметри.",
        "",
        "## Обрана територія",
        "Кампус Національного аерокосмічного університету «ХАІ» та гуртожитки №10-12 у місті Харків, Україна.",
        "",
        "## Обладнання",
        "- PtP: Ubiquiti airMAX LiteBeam 5AC Gen2, 5.15-5.875 ГГц, до 25 dBm, 23 dBi, до 450+ Мбіт/с, довгі лінки до 30+ км.",
        "- PtMP AP: Ubiquiti airMAX Lite AP GPS, 5.15-5.875 ГГц, до 25 dBm, 17 dBi, до 450+ Мбіт/с.",
        "- PtMP Station: Ubiquiti airMAX LiteBeam 5AC Gen2, 5.15-5.875 ГГц, до 25 dBm, 23 dBi, до 450+ Мбіт/с, довгі лінки до 30+ км.",
        "",
        "## Вихідні припущення",
        "- Для профілю траси використано DEM OpenTopoData SRTM90m.",
        "- Частота для всіх радіолінків: 5.5 ГГц, ширина каналу 80 МГц.",
        "- Через відсутність LiDAR у вихідному DEM будівлі та рослинність враховано при виборі висот антен і при текстовому аналізі траси.",
        "",
        "## Завдання 1. Радіолінія PtP",
        f"![PtP profile]({figures['ptp_profiles'].as_posix()})",
        "",
    ]

    for result in results_ptp:
        lines.extend(format_link_result(result))

    lines.extend(
        [
            "## Завдання 2. Радіолінія PtMP",
            f"![PtMP scheme]({figures['ptmp_scheme'].as_posix()})",
            "",
            f"![PtMP profiles]({figures['ptmp_profiles'].as_posix()})",
            "",
        ]
    )
    for result in results_ptmp:
        lines.extend(format_link_result(result))

    lines.extend(
        [
            "## Висновки",
            "- Для траси між головним корпусом ХАІ та гуртожитком №10 при висоті приймальної антени 3 м забезпечується лише nLOS-режим: пряма видимість ще зберігається, але перша зона Френеля перекривається.",
            "- Підняття приймальної антени до 12 м переводить трасу до режиму LOS і зменшує ризик деградації швидкості в міському середовищі.",
            "- Для PtMP-вузла на даху головного корпусу ХАІ три лінки до гуртожитків №10-12 є трасами LOS і забезпечують стійкий рівень сигналу та високу фізичну швидкість навіть з урахуванням міського шуму.",
        ]
    )

    report_path.write_text("\n".join(lines), encoding="utf-8")


def set_default_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)


def add_title_page(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(
        "МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ\n"
        "Національний аерокосмічний університет\n"
        "«Харківський авіаційний інститут»\n\n"
        "Лабораторна робота 7\n"
        "з дисципліни «Антенні пристрої»\n"
        "на тему: «Побудова профілю траси та визначення її основних параметрів»\n"
    )
    run.font.size = Pt(14)

    document.add_paragraph("")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(
        "Виконав: студент 2 курсу групи 526ст/526ст2\n"
        "спеціальності 172 «Телекомунікації та радіотехніка»\n"
        "Чернецький Денис Олександрович\n\n"
        "Перевірила: доцент каф.504\n"
        "Надія Кожемякіна\n"
    )

    document.add_paragraph("\n\nХарків - 2026").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_section(WD_SECTION.NEW_PAGE)


def add_equipment_table(document: Document) -> None:
    document.add_heading("Обладнання", level=1)
    table = document.add_table(rows=1, cols=5)
    header = table.rows[0].cells
    header[0].text = "Пристрій"
    header[1].text = "Тип"
    header[2].text = "Макс. дальність"
    header[3].text = "Макс. ідеальна швидкість"
    header[4].text = "Частота"

    rows = [
        ("Ubiquiti airMAX LiteBeam 5AC Gen2", "PtP / Station PtMP", "30+ км", "450+ Мбіт/с", "5.15-5.875 ГГц"),
        ("Ubiquiti airMAX Lite AP GPS", "AP PtMP", "30+ км", "450+ Мбіт/с", "5.15-5.875 ГГц"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value


def add_result_table(document: Document, result: dict) -> None:
    link = result["link"]
    document.add_heading(link.name, level=2)
    document.add_paragraph(
        f"Довжина лінка: {result['distance_m']:.1f} м ({result['distance_km']:.3f} км). "
        f"Тип траси: {result['link_type']}. Перешкоди: {link.obstacle_type}."
    )
    document.add_paragraph(
        f"Висота антен: передавач {link.tx_height_m:.1f} м, приймач {link.rx_height_m:.1f} м. "
        f"Максимальна перешкода розташована на відстані {result['max_obstacle_distance_m']:.1f} м від передавача."
    )
    document.add_paragraph(
        f"Просвіт LOS у точці максимальної перешкоди: {result['los_clearance_m']:.2f} м. "
        f"1-а зона Френеля у точці максимальної перешкоди: {result['fresnel_at_obstacle_m']:.2f} м. "
        f"На середині траси: {result['fresnel_mid_m']:.2f} м."
    )
    document.add_paragraph(link.note)

    table = document.add_table(rows=1, cols=4)
    header = table.rows[0].cells
    header[0].text = "Сценарій"
    header[1].text = "Expected Signal, dBm"
    header[2].text = "Capacity, Мбіт/с"
    header[3].text = "Модуляція"

    for row in result["radio_rows"]:
        cells = table.add_row().cells
        cells[0].text = row["scenario"]
        cells[1].text = f"{row['signal_dbm']:.1f}"
        cells[2].text = f"{row['capacity_mbps']:.1f}"
        cells[3].text = row["mcs"]


def build_docx_report(results_ptp: list[dict], results_ptmp: list[dict], figures: dict[str, Path]) -> None:
    document = Document()
    set_default_style(document)
    for section in document.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    add_title_page(document)
    document.add_heading("Мета роботи", level=1)
    document.add_paragraph(
        "Закріпити на практиці поняття зони Френеля та профілю траси, "
        "побудувати радіолінії PtP і PtMP на території України та визначити їх основні параметри."
    )

    document.add_heading("Вихідні дані", level=1)
    document.add_paragraph(
        "Обрана територія: кампус Національного аерокосмічного університету «ХАІ» "
        "та гуртожитки №10-12 у місті Харків. Для профілів висот використано DEM "
        "OpenTopoData SRTM90m. Частота для всіх радіолінків - 5.5 ГГц, ширина каналу - 80 МГц."
    )
    document.add_paragraph(
        "Оскільки DEM не містить LiDAR-інформації по будівлях і зелених насадженнях, "
        "висоти антен обрано з запасом для міської забудови."
    )

    add_equipment_table(document)

    document.add_heading("Завдання 1. Радіолінія PtP", level=1)
    document.add_picture(str(figures["ptp_scheme"]), width=Cm(15.5))
    document.add_picture(str(figures["ptp_profiles"]), width=Cm(16.5))
    for result in results_ptp:
        add_result_table(document, result)

    document.add_heading("Завдання 2. Радіолінія PtMP", level=1)
    document.add_picture(str(figures["ptmp_scheme"]), width=Cm(15.5))
    document.add_picture(str(figures["ptmp_profiles"]), width=Cm(16.5))
    for result in results_ptmp:
        add_result_table(document, result)

    document.add_heading("Висновки", level=1)
    document.add_paragraph(
        "Для траси між головним корпусом ХАІ та гуртожитком №10 при висоті приймача 3 м "
        "забезпечується лише режим nLOS: пряма видимість ще є, але перша зона Френеля перекривається."
    )
    document.add_paragraph(
        "Підняття приймальної антени до 12 м переводить трасу у режим LOS та покращує запас "
        "за першою зоною Френеля."
    )
    document.add_paragraph(
        "Для PtMP-вузла на головному корпусі всі три радіолінки до гуртожитків №10-12 є LOS "
        "і відповідають вимогам до рівня сигналу та фізичної швидкості."
    )

    document.save("звіт з лабораторної роботи 7.docx")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    ptp_results = [analyze_link(PTP_NLOS), analyze_link(PTP_LOS)]
    ptmp_results = [analyze_link(link) for link in PTMP_LINKS]
    all_results = ptp_results + ptmp_results

    figures = save_profile_figures(ptp_results, ptmp_results)
    save_csv_outputs(all_results)
    build_markdown_report(ptp_results, ptmp_results, figures)
    build_docx_report(ptp_results, ptmp_results, figures)

    if SHOW_PLOTS:
        plt.show()


if __name__ == "__main__":
    main()
